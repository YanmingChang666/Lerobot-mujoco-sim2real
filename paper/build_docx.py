# -*- coding: utf-8 -*-
"""Build Chinese Word paper + sync key fixes into English Word."""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
MD = ROOT / "全文_双线性Koopman_MPC多机械臂书写.md"
CN_DOCX = ROOT / "koopman_mpc_多机械臂书写_中文.docx"
EN_DOCX = ROOT / "koopman_mpc_soarm101.docx"
FIGS = {
    "figures/x02_center_scene.png": "图 1　X02 居中书写场景。红：参考；绿：实际；浅色：非落笔段。",
    "figures/x02_center_rotated_tracking.png": "图 2　X02 居中、笔姿态 90°、字形 −90° 的“永”字基线结果。",
    "figures/wheelarm_mpc_tracking.png": "图 3　WheelArmR 默认工作区 DBKN–MPC。",
    "figures/x02_mpc_tracking.png": "图 4　X02 默认右侧工作区 DBKN–MPC。",
}


def set_run_font(run, east_asia="宋体", ascii_font="Times New Roman", size=12, bold=False):
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = ascii_font
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.get_or_add_rFonts()
    r_fonts.set(qn("w:ascii"), ascii_font)
    r_fonts.set(qn("w:hAnsi"), ascii_font)
    r_fonts.set(qn("w:eastAsia"), east_asia)


def set_paragraph_format(p, first_line=True, space_after=6, line=1.5, align=None, space_before=0):
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing = line
    pf.first_line_indent = Cm(0.74) if first_line else Cm(0)
    if align is not None:
        p.alignment = align


def strip_md_inline(text: str) -> str:
    text = text.replace("**", "")
    text = re.sub(r"`([^`]+)`", r"\1", text)
    text = re.sub(r"!?\[([^\]]*)\]\([^)]+\)", r"\1", text)
    text = text.replace(r"\,", " ")
    return text


def add_mixed_paragraph(doc, text, style="body"):
    text = strip_md_inline(text).strip()
    if not text:
        return None
    p = doc.add_paragraph()
    if style == "title":
        set_paragraph_format(p, first_line=False, space_after=12, line=1.3, align=WD_ALIGN_PARAGRAPH.CENTER)
        run = p.add_run(text)
        set_run_font(run, east_asia="黑体", size=16, bold=True)
    elif style == "h1":
        set_paragraph_format(p, first_line=False, space_after=8, space_before=14)
        run = p.add_run(text)
        set_run_font(run, east_asia="黑体", size=14, bold=True)
    elif style == "h2":
        set_paragraph_format(p, first_line=False, space_after=6, space_before=10)
        run = p.add_run(text)
        set_run_font(run, east_asia="黑体", size=12, bold=True)
    elif style == "caption":
        set_paragraph_format(p, first_line=False, space_after=10, align=WD_ALIGN_PARAGRAPH.CENTER)
        run = p.add_run(text)
        set_run_font(run, east_asia="宋体", size=10.5)
    elif style == "keywords":
        set_paragraph_format(p, first_line=False, space_after=10)
        run = p.add_run(text)
        set_run_font(run, east_asia="宋体", size=10.5)
    else:
        set_paragraph_format(p, first_line=True, space_after=6)
        run = p.add_run(text)
        set_run_font(run, east_asia="宋体", size=12)
    return p


def add_table_from_rows(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    for j, h in enumerate(headers):
        cell = table.rows[0].cells[j]
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(h)
        set_run_font(run, east_asia="宋体", size=9, bold=True)
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = table.rows[i + 1].cells[j]
            cell.text = ""
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(val)
            set_run_font(run, east_asia="宋体", size=9)
    doc.add_paragraph()


def parse_md_table(lines, start_idx):
    rows_raw = []
    i = start_idx
    while i < len(lines) and lines[i].strip().startswith("|"):
        rows_raw.append(lines[i].strip())
        i += 1

    def split_row(s):
        parts = [c.strip() for c in s.strip("|").split("|")]
        return [strip_md_inline(p) for p in parts]

    if len(rows_raw) < 2:
        return None, None, start_idx + 1
    headers = split_row(rows_raw[0])
    data = [split_row(r) for r in rows_raw[2:]]
    n = len(headers)
    data = [(r + [""] * (n - len(r)))[:n] for r in data]
    return headers, data, i


def simplify_latex(formula: str) -> str:
    formula = formula.replace(r"\Bigl", "").replace(r"\Bigr", "")
    formula = formula.replace(r"\bigl", "").replace(r"\bigr", "")
    formula = re.sub(r"\\[a-zA-Z]+\{([^}]*)\}", r"\1", formula)
    formula = formula.replace("\\", "")
    return formula


def build_chinese_docx():
    lines = MD.read_text(encoding="utf-8").splitlines()
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)

    i = 0
    title_done = False
    while i < len(lines):
        s = lines[i].strip()
        if not s or s == "---" or s.startswith(">"):
            i += 1
            continue

        if s.startswith("# ") and not title_done:
            add_mixed_paragraph(doc, s[2:], style="title")
            title_done = True
            i += 1
            continue
        if s.startswith("## "):
            add_mixed_paragraph(doc, s[3:], style="h1")
            i += 1
            continue
        if s.startswith("### ") or s.startswith("#### "):
            add_mixed_paragraph(doc, s.lstrip("#").strip(), style="h2")
            i += 1
            continue

        m = re.match(r"!\[([^\]]*)\]\(([^)]+)\)", s)
        if m:
            path = m.group(2)
            img = ROOT / path
            if img.exists():
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.first_line_indent = Cm(0)
                p.add_run().add_picture(str(img), width=Cm(14.0))
            cap = FIGS.get(path)
            j = i + 1
            while j < len(lines) and not lines[j].strip():
                j += 1
            if j < len(lines) and ("图 " in lines[j] or lines[j].strip().startswith("**图")):
                cap = strip_md_inline(lines[j].strip())
                i = j + 1
            else:
                i += 1
            if cap:
                add_mixed_paragraph(doc, cap, style="caption")
            continue

        if s.startswith("$$"):
            if s == "$$":
                i += 1
                buf = []
                while i < len(lines) and lines[i].strip() != "$$":
                    buf.append(lines[i].strip())
                    i += 1
                i += 1
                formula = " ".join(buf)
            else:
                formula = s.strip("$")
                i += 1
            p = doc.add_paragraph()
            set_paragraph_format(p, first_line=False, align=WD_ALIGN_PARAGRAPH.CENTER)
            run = p.add_run(simplify_latex(formula))
            set_run_font(run, ascii_font="Times New Roman", east_asia="宋体", size=11)
            continue

        if s.startswith("|"):
            headers, rows, ni = parse_md_table(lines, i)
            if headers:
                add_table_from_rows(doc, headers, rows)
                i = ni
                continue

        if re.match(r"^\d+\.\s+", s) or s.startswith("- "):
            content = re.sub(r"^\d+\.\s+", "", s)
            if content.startswith("- "):
                content = content[2:]
            p = doc.add_paragraph()
            set_paragraph_format(p, first_line=False, space_after=4)
            p.paragraph_format.left_indent = Cm(0.74)
            run = p.add_run(strip_md_inline(content))
            set_run_font(run, size=12)
            i += 1
            continue

        if s.startswith("```"):
            i += 1
            buf = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                buf.append(lines[i])
                i += 1
            i += 1
            p = doc.add_paragraph()
            set_paragraph_format(p, first_line=False, space_after=6, line=1.15)
            run = p.add_run("\n".join(buf))
            set_run_font(run, ascii_font="Consolas", east_asia="宋体", size=9)
            continue

        if s.startswith("**关键词") or s.startswith("关键词"):
            add_mixed_paragraph(doc, strip_md_inline(s), style="keywords")
            i += 1
            continue

        add_mixed_paragraph(doc, s, style="body")
        i += 1

    p = doc.add_paragraph()
    set_paragraph_format(p, first_line=False, space_after=0)
    run = p.add_run(
        "（本稿由 paper/全文_双线性Koopman_MPC多机械臂书写.md 生成；"
        "实验快照见 paper/evidence/，可用 python paper/verify_results.py 复算。）"
    )
    set_run_font(run, east_asia="宋体", size=9)
    run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    doc.save(CN_DOCX)
    print(f"Wrote {CN_DOCX} ({CN_DOCX.stat().st_size} bytes)")


def replace_in_paragraph(paragraph, old, new):
    if old not in paragraph.text:
        return False
    new_full = paragraph.text.replace(old, new)
    if not paragraph.runs:
        paragraph.add_run(new_full)
        return True
    paragraph.runs[0].text = new_full
    for r in paragraph.runs[1:]:
        r.text = ""
    return True


def sync_english_docx():
    backup = ROOT / "koopman_mpc_soarm101.docx.bak"
    if not backup.exists():
        backup.write_bytes(EN_DOCX.read_bytes())
        print(f"Backup -> {backup}")

    doc = Document(str(EN_DOCX))
    replacements = [
        (
            "that KESO reduces payload-induced tracking error by roughly 56× and combined noise+payload error by roughly 327× relative to an uncompensated bilinear Koopman-MPC baseline",
            "that KESO reduces payload-induced tracking error by roughly 17.5× relative to the UKF-only payload condition (0.263 → 0.015) and combined noise+payload error by roughly 330× relative to a fully uncompensated bilinear Koopman-MPC baseline",
        ),
        (
            "KESO reduces this back down by ~56× relative to the uncompensated-with-UKF payload condition (0.263 → 0.0150, essentially recovering baseline-level tracking), and the combined noise+payload condition is reduced by ~327× relative to the fully uncompensated case (4.967 → 0.0152)",
            "KESO reduces this back down by ~17.5× relative to the UKF-only payload condition (0.263 → 0.015, essentially recovering baseline-level tracking), and the combined noise+payload condition is reduced by ~330× relative to the fully uncompensated case (4.967 → 0.015)",
        ),
        (
            "a spectral-radius penalty discouraging unstable roll-out of A (eigenvalues held below 0.99 in magnitude), and an L1 penalty on H discouraging spurious state-control coupling. For non-invertible architectures an additional autoencoder reconstruction term is included; for IKN/IBKN this term is identically zero by construction and is omitted.",
            "a soft spectral-radius penalty on A (discouraging eigenvalues with magnitude above 0.99; not a hard stability certificate for the controlled bilinear system). An L1 penalty on H is logged in the current codebase but is not added to the backpropagated total loss. For DKUC/DBKN the state is recovered by a fixed projection (no trained decoder); a reconstruction term applies only when a trainable decoder is enabled. For IKN/IBKN the inverse is available in closed form and no reconstruction loss is used.",
        ),
        (
            "in simulation (MuJoCo [9]) and on the physical arm:",
            "in simulation (MuJoCo [10]) and on the physical arm:",
        ),
        (
            "an ADRC-style total-disturbance estimator [8] relocated from joint space",
            "an ADRC-style total-disturbance estimator [9] relocated from joint space",
        ),
        (
            "reproduces arbitrary Chinese characters from public stroke-order data [11].",
            "reproduces multi-stroke Chinese characters from public stroke-order data [12], within the reachable workspace covered by training.",
        ),
        (
            "The Unscented Kalman Filter [8] is the standard nonlinear extension of the Kalman filter; we are not aware of prior work applying it directly inside a learned Koopman-lifted state rather than the physical state. Active Disturbance Rejection Control and its Extended State Observer [8] are classically formulated",
            "The Unscented Kalman Filter [8] is the standard nonlinear extension of the Kalman filter; placing it in learned Koopman-lifted coordinates that share A, B, H with MPC is an engineering choice of this study (priority relative to other lifted-space filters should be audited separately; we do not claim absolute novelty). Active Disturbance Rejection Control and its Extended State Observer [9] are classically formulated",
        ),
        (
            "is solved with CasADi [10];",
            "is solved with CasADi [11];",
        ),
        (
            "an Extended State Observer in the spirit of ADRC [8], but operating",
            "an Extended State Observer in the spirit of ADRC [9], but operating",
        ),
        (
            "Chinese characters sourced from public stroke-order data [11].",
            "Chinese characters sourced from public stroke-order data [12].",
        ),
        (
            "that reproduces arbitrary Chinese-character stroke data",
            "that reproduces multi-stroke Chinese-character stroke data within reachable, training-covered workspace",
        ),
    ]

    n_rep = 0
    for p in doc.paragraphs:
        for old, new in replacements:
            if replace_in_paragraph(p, old, new):
                n_rep += 1
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for old, new in replacements:
                        if replace_in_paragraph(p, old, new):
                            n_rep += 1

    joined = "\n".join(p.text for p in doc.paragraphs)
    if "Note on Chinese full manuscript" not in joined:
        doc.add_paragraph()
        p = doc.add_paragraph()
        p.add_run(
            "Note on Chinese full manuscript / 中文完整稿. "
            "A merged Chinese manuscript that also covers multi-arm handwriting "
            "(Piper / WheelArmR / X02), YAML-based robot onboarding, corrected "
            "disturbance-rejection ratios (~17.5× for 0.263→0.015), loss-function "
            "wording aligned with the repository, and citation numbering "
            "(UKF [8], ADRC [9], MuJoCo [10], CasADi [11], stroke data [12]) "
            "is provided as paper/koopman_mpc_多机械臂书写_中文.docx "
            "(source: paper/全文_双线性Koopman_MPC多机械臂书写.md)."
        )

    doc.save(EN_DOCX)
    print(f"Updated {EN_DOCX}: {n_rep} paragraph replacements")


if __name__ == "__main__":
    build_chinese_docx()
    sync_english_docx()
    print("Done.")
